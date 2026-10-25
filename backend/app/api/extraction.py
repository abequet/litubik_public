# SPDX-FileCopyrightText: 2026 Université Gustave Eiffel
# SPDX-License-Identifier: AGPL-3.0-or-later
"""M9: data extraction — config, dynamic columns, grid values, study-design
qualification, JBI critical appraisal, XLSX/CSV export."""
import csv
import io
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

from app.api.deps import (
    DbDep, ProjectMemberDep, SatelliteGuard, get_project,
    require_project_admin,
)
from app.core.csvsafe import safe_row
from app.models import (
    ExtractionColumn, ExtractionConfig, ExtractionStudy, ExtractionValue,
    JbiAppraisal, ProjectMember, Record, User,
)
from app.services import framing
from app.services.extraction import ai as ext_ai
from app.services.appraisal import core as appraisal_core
from app.services.appraisal import instruments
from app.services.extraction import service as ext
from app.services.extraction import studies as ext_studies
from app.services.llm.base import LLMError

router = APIRouter(prefix="/api/projects/{project_id}/extraction",
                   tags=["extraction"])


def _record(project_id: int, record_id: int, db) -> Record:
    record = db.get(Record, record_id)
    if record is None or record.project_id != project_id:
        raise HTTPException(status_code=404, detail="Record not found")
    return record


def _study(project_id: int, db, record_id: int | None = None,
           study_id: int | None = None, user_id: int | None = None) -> ExtractionStudy:
    """The study a request targets: explicit study_id, or the default study
    of record_id (chantier 9 — pre-study clients keep working)."""
    try:
        return ext_studies.resolve_study(db, project_id, record_id, study_id, user_id)
    except LookupError as e:
        raise HTTPException(status_code=404, detail=str(e))


# ---- config -----------------------------------------------------------------------

class ConfigIn(BaseModel):
    n_extractors: int = Field(default=1, ge=1, le=2)
    extractor_ids: list[int] = []


@router.get("/config")
def get_config(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    config = ext.get_config(db, project_id)
    members = (db.query(ProjectMember, User)
               .join(User, User.id == ProjectMember.user_id)
               .filter(ProjectMember.project_id == project_id).all())
    return {
        "config": None if config is None else {
            "n_extractors": config.n_extractors,
            "extractor_ids": config.extractor_ids or [],
            "status": config.status,
            "started_at": config.started_at.isoformat() if config.started_at else None,
        },
        "members": [{"user_id": m.user_id, "display_name": u.display_name,
                     "role": m.role} for m, u in members],
        "eligible_records": len(ext.eligible_ids(db, project_id)),
        "me": member.user_id,
        "my_role": member.role,
    }


@router.put("/config",
            dependencies=[Depends(require_project_admin), SatelliteGuard])
def put_config(project_id: int, payload: ConfigIn, member: ProjectMemberDep,
               db: DbDep):
    get_project(project_id, db)
    member_ids = {m.user_id for m in db.query(ProjectMember)
                  .filter_by(project_id=project_id)}
    bad = [u for u in payload.extractor_ids if u not in member_ids]
    if bad:
        raise HTTPException(status_code=400,
                            detail="Extractors must be project members")
    config = ext.get_config(db, project_id)
    if config is None:
        config = ExtractionConfig(project_id=project_id)
        db.add(config)
    config.n_extractors = payload.n_extractors
    config.extractor_ids = payload.extractor_ids
    db.commit()
    return {"ok": True}


@router.post("/start",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
def start(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    config = ext.get_config(db, project_id)
    if config is None or not config.extractor_ids:
        raise HTTPException(status_code=400,
                            detail="Choose the extractors first")
    if not ext.eligible_ids(db, project_id):
        raise HTTPException(status_code=400,
                            detail="No full-text-included record to extract "
                                   "yet — finish full-text screening first")
    created = ext.start(db, config)
    return {"ok": True, "assignments_created": created}


# ---- columns ----------------------------------------------------------------------

class ColumnIn(BaseModel):
    name: str = Field(min_length=1, max_length=255)
    col_type: str = Field(default="text",
                          pattern="^(text|number|bool|choice|multi)$")
    description: str = ""
    options: list[str] = []
    position: int | None = None


@router.get("/columns")
def list_columns(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    columns = (db.query(ExtractionColumn).filter_by(project_id=project_id)
               .order_by(ExtractionColumn.position, ExtractionColumn.id).all())
    return [{"id": c.id, "name": c.name, "col_type": c.col_type,
             "description": c.description, "options": c.options or [],
             "position": c.position} for c in columns]


@router.post("/columns",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
def add_column(project_id: int, payload: ColumnIn, member: ProjectMemberDep,
               db: DbDep):
    get_project(project_id, db)
    if payload.col_type in ("choice", "multi") and not payload.options:
        raise HTTPException(status_code=400,
                            detail="choice/multi columns need options")
    max_pos = (db.query(ExtractionColumn)
               .filter_by(project_id=project_id).count())
    column = ExtractionColumn(
        project_id=project_id, name=payload.name, col_type=payload.col_type,
        description=payload.description, options=payload.options,
        position=payload.position if payload.position is not None else max_pos)
    db.add(column)
    db.commit()
    return {"id": column.id}


@router.put("/columns/{column_id}",
            dependencies=[Depends(require_project_admin), SatelliteGuard])
def update_column(project_id: int, column_id: int, payload: ColumnIn,
                  member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    column = db.get(ExtractionColumn, column_id)
    if column is None or column.project_id != project_id:
        raise HTTPException(status_code=404, detail="Column not found")
    column.name = payload.name
    column.col_type = payload.col_type
    column.description = payload.description
    column.options = payload.options
    if payload.position is not None:
        column.position = payload.position
    db.commit()
    return {"ok": True}


@router.delete("/columns/{column_id}",
               dependencies=[Depends(require_project_admin), SatelliteGuard])
def delete_column(project_id: int, column_id: int, member: ProjectMemberDep,
                  db: DbDep):
    get_project(project_id, db)
    column = db.get(ExtractionColumn, column_id)
    if column is None or column.project_id != project_id:
        raise HTTPException(status_code=404, detail="Column not found")
    db.delete(column)
    db.commit()
    return {"ok": True}


@router.get("/columns/templates")
def list_templates(project_id: int, member: ProjectMemberDep, db: DbDep):
    """The starter column sets, with the one this review's type suggests
    flagged — the others stay pickable."""
    project = get_project(project_id, db)
    default = framing.review_type(project.review_type).default_template
    return [{"key": t.key, "label": t.label, "description": t.description,
             "n_columns": len(t.columns), "is_default": t.key == default}
            for t in ext.TEMPLATES.values()]


class TemplateIn(BaseModel):
    template: str | None = None


@router.post("/columns/template",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
def add_template(project_id: int, member: ProjectMemberDep, db: DbDep,
                 payload: TemplateIn | None = None):
    """One-click starter columns (skips names already present).

    With no body, the template of the project's review type is used — which is
    the standard evidence table for a plain systematic review, i.e. exactly
    what this endpoint did before templates were a catalogue.
    """
    project = get_project(project_id, db)
    key = (payload.template if payload and payload.template
           else framing.review_type(project.review_type).default_template)
    template = ext.TEMPLATES.get(key)
    if template is None:
        raise HTTPException(status_code=400, detail=f"Unknown template: {key}")
    existing = {c.name.lower() for c in db.query(ExtractionColumn)
                .filter_by(project_id=project_id)}
    pos = len(existing)
    created = 0
    for template_column in template.columns:
        if template_column["name"].lower() in existing:
            continue
        db.add(ExtractionColumn(project_id=project_id, position=pos,
                                **template_column))
        pos += 1
        created += 1
    db.commit()
    return {"created": created, "template": template.key}


@router.post("/columns/suggest",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
def suggest_columns(project_id: int, member: ProjectMemberDep, db: DbDep):
    """AI column proposal from the validated protocol — returns suggestions,
    the admin picks which to add."""
    get_project(project_id, db)
    try:
        return ext_ai.suggest_columns(db, project_id, user_id=member.user_id)
    except LLMError as e:
        raise HTTPException(status_code=400, detail=str(e))


# ---- external table import --------------------------------------------------------

MAX_IMPORT_BYTES = 20 * 1024 * 1024


@router.post("/import-table/preview",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
async def import_table_preview(project_id: int, file: UploadFile,
                               member: ProjectMemberDep, db: DbDep):
    """Parse an extraction table filled on another platform (xlsx): proposed
    header mapping + row-matching stats. Nothing is written."""
    get_project(project_id, db)
    data = await file.read()
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 20 MB)")
    from app.services.extraction import import_table
    try:
        return import_table.preview(db, project_id, data)
    except import_table.ImportTableError as e:
        raise HTTPException(status_code=400, detail=str(e))


@router.post("/import-table",
             dependencies=[Depends(require_project_admin), SatelliteGuard])
async def import_table_apply(project_id: int, file: UploadFile,
                             member: ProjectMemberDep, db: DbDep,
                             mapping: str = Form(...),
                             create_missing: bool = Form(default=True),
                             source_name: str = Form(
                                 default="Extraction import", min_length=1,
                                 max_length=100)):
    """Apply the import with the admin-adjusted mapping: create the missing
    columns/records, force-include the imported articles into the extraction
    corpus and store the non-empty cells as the importer's validated values."""
    get_project(project_id, db)
    data = await file.read()
    if len(data) > MAX_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 20 MB)")
    from app.services.extraction import import_table
    safe_name = (file.filename or "extraction.xlsx") \
        .replace("\\", "/").rsplit("/", 1)[-1]
    try:
        result = import_table.apply(
            db, project_id, member.user_id, safe_name, data,
            mapping_raw=mapping, create_missing=create_missing,
            source_name=source_name)
    except import_table.ImportTableError as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=str(e))
    db.commit()
    return result


# ---- grid + values ----------------------------------------------------------------

@router.get("/table")
def table(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    return ext.table_data(db, project_id, me=member.user_id)


class ValueIn(BaseModel):
    record_id: int | None = None   # default study of the article…
    study_id: int | None = None    # …or one study explicitly
    column_id: int
    value: object = None
    validated: bool = True  # a human save is a validation by default


@router.put("/values")
def put_value(project_id: int, payload: ValueIn, member: ProjectMemberDep,
              db: DbDep):
    get_project(project_id, db)
    study = _study(project_id, db, payload.record_id, payload.study_id,
                   member.user_id)
    column = db.get(ExtractionColumn, payload.column_id)
    if column is None or column.project_id != project_id:
        raise HTTPException(status_code=404, detail="Column not found")
    try:
        stored = ext.coerce_value(column, payload.value)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    value = (db.query(ExtractionValue)
             .filter_by(study_id=study.id, column_id=column.id,
                        user_id=member.user_id).first())
    if value is None:
        value = ExtractionValue(record_id=study.record_id, study_id=study.id,
                                column_id=column.id, user_id=member.user_id)
        db.add(value)
    value.value = stored
    value.origin = "human"
    value.validated = payload.validated
    db.commit()
    return {"ok": True}


class ValidateIn(BaseModel):
    record_id: int | None = None
    study_id: int | None = None
    column_id: int


@router.post("/values/validate")
def validate_ai_value(project_id: int, payload: ValidateIn,
                      member: ProjectMemberDep, db: DbDep):
    """Accept an AI-proposed cell as-is (turns it into my validated value)."""
    get_project(project_id, db)
    study = _study(project_id, db, payload.record_id, payload.study_id,
                   member.user_id)
    value = (db.query(ExtractionValue)
             .filter_by(study_id=study.id,
                        column_id=payload.column_id,
                        user_id=member.user_id).first())
    if value is None:
        raise HTTPException(status_code=404, detail="No value to validate")
    value.validated = True
    db.commit()
    return {"ok": True}


@router.post("/values/validate-all")
def validate_all_ai_values(project_id: int, member: ProjectMemberDep, db: DbDep):
    """Accept every AI-proposed cell of mine that is still awaiting validation
    (the 'validate all' action on the extraction table)."""
    get_project(project_id, db)
    pending = (db.query(ExtractionValue)
               .join(Record, Record.id == ExtractionValue.record_id)
               .filter(Record.project_id == project_id,
                       ExtractionValue.user_id == member.user_id,
                       ExtractionValue.origin == "ai",
                       ExtractionValue.validated.is_(False))
               .all())
    for value in pending:
        value.validated = True
    db.commit()
    return {"validated": len(pending)}


@router.get("/records/{record_id}/ai-estimate")
def record_ai_estimate(project_id: int, record_id: int, member: ProjectMemberDep,
                       db: DbDep, kind: str = "prefill"):
    """Up-front cost estimate for a per-article AI button (fill / design / jbi)."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    kind = kind if kind in ("prefill", "design", "jbi") else "prefill"
    # prefill and appraisal read the PDF: refused up-front when it is over the
    # page limit (the design suggestion falls back to title + abstract)
    blocked = (ext_ai.pdf_block_reason(db, record)
               if kind in ("prefill", "jbi") else None)
    return {"cost": (None if blocked else
                     ext_ai.estimate_record_cost(db, project_id, record, kind)),
            "blocked": blocked}


@router.post("/records/{record_id}/prefill")
def prefill(project_id: int, record_id: int, member: ProjectMemberDep,
            db: DbDep):
    """On-demand AI prefill of one record's row from its stored PDF."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    try:
        return ext_ai.prefill_record(db, project_id, record,
                                     user_id=member.user_id)
    except LLMError as e:
        raise HTTPException(status_code=400, detail=str(e))


# ---- study design -----------------------------------------------------------------

class DesignIn(BaseModel):
    design: str = Field(pattern="^([a-z_]+|)$")


def _apply_design(db, study: ExtractionStudy, design: str) -> None:
    # The design vocabulary is shared across instruments — a study's design is
    # a property of the study, not of the tool used to appraise it.
    if design not in (set(instruments.DESIGNS) | {""}):
        raise HTTPException(status_code=400, detail="Unknown design")
    ext_studies.set_design(db, study, design or None, "human" if design else None)
    db.commit()


@router.put("/records/{record_id}/design", dependencies=[SatelliteGuard])
def set_design(project_id: int, record_id: int, payload: DesignIn,
               member: ProjectMemberDep, db: DbDep):
    """Design of the article's first (default) study."""
    get_project(project_id, db)
    study = _study(project_id, db, record_id=record_id, user_id=member.user_id)
    _apply_design(db, study, payload.design)
    return {"ok": True}


@router.put("/studies/{study_id}/design", dependencies=[SatelliteGuard])
def set_study_design(project_id: int, study_id: int, payload: DesignIn,
                     member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    study = _study(project_id, db, study_id=study_id)
    _apply_design(db, study, payload.design)
    return {"ok": True}


# ---- studies within an article (chantier 9) ----------------------------------------

class StudyIn(BaseModel):
    label: str = Field(default="", max_length=120)


@router.get("/records/{record_id}/studies")
def list_studies(project_id: int, record_id: int, member: ProjectMemberDep,
                 db: DbDep):
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    studies = ext_studies.studies_for_records(db, [record.id], member.user_id)
    return {"studies": [ext_studies.study_dict(s) for s in studies[record.id]]}


@router.post("/records/{record_id}/studies")
def add_study(project_id: int, record_id: int, payload: StudyIn,
              member: ProjectMemberDep, db: DbDep):
    """Declare another study reported by this article. Any member may: the
    extractor reading the paper is the one who finds out."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    try:
        study = ext_studies.add_study(db, record, payload.label, member.user_id)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    db.commit()
    return ext_studies.study_dict(study)


@router.put("/studies/{study_id}")
def rename_study(project_id: int, study_id: int, payload: StudyIn,
                 member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    study = _study(project_id, db, study_id=study_id)
    ext_studies.rename_study(db, study, payload.label)
    db.commit()
    return ext_studies.study_dict(study)


@router.delete("/studies/{study_id}")
def delete_study(project_id: int, study_id: int, member: ProjectMemberDep,
                 db: DbDep):
    """Remove a study and its values/appraisals (the first study of an
    article stays)."""
    get_project(project_id, db)
    study = _study(project_id, db, study_id=study_id)
    try:
        removed = ext_studies.delete_study(db, study)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    db.commit()
    return {"ok": True, **removed}


@router.post("/records/{record_id}/design/suggest",
             dependencies=[SatelliteGuard])
def suggest_design(project_id: int, record_id: int, member: ProjectMemberDep,
                   db: DbDep):
    """AI qualification of the study design (per article, on demand). The
    suggestion is stored (origin=ai) but stays overridable by hand."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    try:
        suggestion = ext_ai.suggest_design(db, project_id, record,
                                           user_id=member.user_id)
    except LLMError as e:
        raise HTTPException(status_code=400, detail=str(e))
    # the AI reads the whole paper: its call qualifies the first study, and
    # never overrides a design a human set
    study = _study(project_id, db, record_id=record.id, user_id=member.user_id)
    if study.design_origin != "human" and suggestion["design"] != "other":
        ext_studies.set_design(db, study, suggestion["design"], "ai")
        db.commit()
    return suggestion


# ---- JBI critical appraisal --------------------------------------------------------

@router.get("/jbi/checklists")
def checklists(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    # The checklists of the instrument this project appraises with — JBI
    # unless it declared its own (see /appraisal/instruments).
    inst = appraisal_core.instrument_of(db, project_id)
    catalogue = instruments.as_catalogue(inst)
    return {"checklists": catalogue["checklists"],
            "instrument": instruments.summary(inst),
            "overall_options": instruments.OVERALL_OPTIONS}


@router.post("/records/{record_id}/jbi/suggest")
def suggest_jbi(project_id: int, record_id: int, member: ProjectMemberDep,
                db: DbDep):
    """On-demand AI suggestions for the record's JBI checklist, read from the
    stored PDF. Returns per-question {answer, comment} proposals — NOTHING is
    saved; the assessor reviews the prefill and saves their own appraisal."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    try:
        return ext_ai.suggest_jbi(db, project_id, record,
                                  user_id=member.user_id)
    except LLMError as e:
        raise HTTPException(status_code=400, detail=str(e))


class AppraisalIn(BaseModel):
    answers: dict = {}
    overall: str = Field(default="", pattern="^(include|exclude|seek_info|)$")
    comment: str = ""


@router.get("/records/{record_id}/jbi")
def get_appraisals(project_id: int, record_id: int, member: ProjectMemberDep,
                   db: DbDep, study_id: int | None = None):
    """My appraisal (editable) + the other assessors' (read-only, named) of
    one study — the article's first one unless study_id says otherwise."""
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    study = _study(project_id, db, record.id, study_id, member.user_id)
    if study.record_id != record.id:
        raise HTTPException(status_code=404, detail="Study not found")
    names = {u.id: u.display_name for u in db.query(User)}
    rows = db.query(JbiAppraisal).filter_by(study_id=study.id).all()

    def out(a: JbiAppraisal) -> dict:
        return {"user_id": a.user_id, "user": names.get(a.user_id, "?"),
                "design": a.design, "answers": a.answers or {},
                "overall": a.overall, "comment": a.comment,
                "updated_at": a.updated_at.isoformat() if a.updated_at else ""}

    return {
        "study": ext_studies.study_dict(study),
        "design": study.study_design, "design_origin": study.design_origin,
        "mine": next((out(a) for a in rows if a.user_id == member.user_id), None),
        "others": [out(a) for a in rows if a.user_id != member.user_id],
    }


@router.put("/records/{record_id}/jbi")
def put_appraisal(project_id: int, record_id: int, payload: AppraisalIn,
                  member: ProjectMemberDep, db: DbDep,
                  study_id: int | None = None):
    get_project(project_id, db)
    record = _record(project_id, record_id, db)
    study = _study(project_id, db, record.id, study_id, member.user_id)
    if study.record_id != record.id:
        raise HTTPException(status_code=404, detail="Study not found")
    if not study.study_design or study.study_design == "other":
        raise HTTPException(status_code=400,
                            detail="Qualify the study design first — it "
                                   "selects the checklist")
    inst = appraisal_core.instrument_of(db, project_id)
    checklist = instruments.checklist_for(inst, study.study_design)
    if checklist is None:
        raise HTTPException(
            status_code=400,
            detail=f"{inst.label} has no checklist for this study design")
    valid_ids = set(instruments.question_ids(inst, study.study_design))
    valid_answers = set(checklist["options"])
    answers = {}
    for qid, entry in (payload.answers or {}).items():
        if qid not in valid_ids or not isinstance(entry, dict):
            continue
        answer = str(entry.get("answer", ""))
        if answer and answer not in valid_answers:
            raise HTTPException(status_code=400,
                                detail=f"Invalid answer for {qid}: {answer}")
        answers[qid] = {"answer": answer,
                        "comment": str(entry.get("comment", ""))}
    appraisal = (db.query(JbiAppraisal)
                 .filter_by(study_id=study.id, user_id=member.user_id)
                 .first())
    if appraisal is None:
        appraisal = JbiAppraisal(record_id=record.id, study_id=study.id,
                                 user_id=member.user_id,
                                 design=study.study_design)
        db.add(appraisal)
    appraisal.design = study.study_design
    appraisal.instrument = inst.key
    appraisal.answers = answers
    # an instrument without an overall judgement (MMAT) records none, whatever
    # a stale form sent: the per-criterion ratings are the appraisal
    appraisal.overall = payload.overall if inst.overall else ""
    appraisal.comment = payload.comment
    db.commit()
    return {"ok": True}


# ---- exports ----------------------------------------------------------------------

@router.get("/export.csv")
def export_csv(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    header, rows = ext.export_rows(db, project_id)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(safe_row(header))
    writer.writerows(safe_row(row) for row in rows)
    return StreamingResponse(
        iter([buffer.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition":
                 'attachment; filename="extraction.csv"'})


@router.get("/export.xlsx")
def export_xlsx(project_id: int, member: ProjectMemberDep, db: DbDep):
    get_project(project_id, db)
    from openpyxl import Workbook
    from openpyxl.styles import Font

    header, rows = ext.export_rows(db, project_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extraction"
    sheet.append(safe_row(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(safe_row(row))
    sheet.freeze_panes = "A2"

    # appraisal sheet: one line per (study, assessor, question)
    detail_header, detail_rows = _appraisal_details(db, project_id)
    if detail_rows:
        jbi_sheet = workbook.create_sheet("Appraisal")
        jbi_sheet.append(safe_row(detail_header))
        for cell in jbi_sheet[1]:
            cell.font = Font(bold=True)
        for line in detail_rows:
            jbi_sheet.append(safe_row(line))

    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type=("application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"),
        headers={"Content-Disposition":
                 'attachment; filename="extraction.xlsx"'})


def _appraisal_details(db, project_id: int) -> tuple[list[str], list[list]]:
    """One line per (study, assessor, question): the question's full text,
    the answer as the instrument spells it, the evidence comment and — when
    the instrument has one — the overall judgement."""
    inst = appraisal_core.instrument_of(db, project_id)
    names = {u.id: u.display_name for u in db.query(User)}
    appraisals = (db.query(JbiAppraisal).join(
        Record, Record.id == JbiAppraisal.record_id)
        .filter(Record.project_id == project_id)
        .order_by(JbiAppraisal.record_id, JbiAppraisal.study_id,
                  JbiAppraisal.user_id).all())
    study_labels = {s.id: s.label for s in db.query(ExtractionStudy)
                    .join(Record, Record.id == ExtractionStudy.record_id)
                    .filter(Record.project_id == project_id)}
    header = ["Record ID", "Study", "Assessor", "Design", "Instrument",
              "Criterion", "Question", "Answer", "Comment"]
    if inst.overall:
        header.append("Overall")
    rows = []
    for appraisal in appraisals:
        checklist = instruments.checklist_for(inst, appraisal.design) or {}
        labels = checklist.get("option_labels") or {}
        questions = {q["id"]: q["text"]
                     for d in checklist.get("domains", [])
                     for q in d["questions"]}
        for qid, entry in (appraisal.answers or {}).items():
            answer = entry.get("answer", "")
            line = [appraisal.record_id,
                    study_labels.get(appraisal.study_id, ""),
                    names.get(appraisal.user_id, "?"),
                    appraisal.design, appraisal.instrument, qid,
                    questions.get(qid, ""), labels.get(answer, answer),
                    entry.get("comment", "")]
            if inst.overall:
                line.append(appraisal.overall)
            rows.append(line)
    return header, rows


def _appraisal_filename(db, project_id: int, ext_: str) -> str:
    inst = appraisal_core.instrument_of(db, project_id)
    return f'attachment; filename="{inst.key}_appraisal.{ext_}"'


@router.get("/export/jbi.csv")
def export_jbi_csv(project_id: int, member: ProjectMemberDep, db: DbDep):
    """Dedicated appraisal table: one line per study × assessor, wide answers
    under the instrument's own criterion ids, plus the score when the
    instrument has one. The route keeps its historical name."""
    get_project(project_id, db)
    header, rows = ext.appraisal_export_rows(db, project_id)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(safe_row(header))
    writer.writerows(safe_row(r) for r in rows)
    return StreamingResponse(
        iter([buffer.getvalue().encode("utf-8-sig")]),
        media_type="text/csv",
        headers={"Content-Disposition":
                 _appraisal_filename(db, project_id, "csv")})


@router.get("/export/jbi.xlsx")
def export_jbi_xlsx(project_id: int, member: ProjectMemberDep, db: DbDep):
    """Dedicated appraisal workbook: the wide table + a details sheet with
    every question's full text, answer and evidence comment."""
    get_project(project_id, db)
    from openpyxl import Workbook
    from openpyxl.styles import Font

    header, rows = ext.appraisal_export_rows(db, project_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Appraisal"
    sheet.append(safe_row(header))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(safe_row(row))
    sheet.freeze_panes = "A2"

    detail_header, detail_rows = _appraisal_details(db, project_id)
    details = workbook.create_sheet("Details")
    details.append(safe_row(detail_header))
    for cell in details[1]:
        cell.font = Font(bold=True)
    for line in detail_rows:
        details.append(safe_row(line))
    details.freeze_panes = "A2"

    out = io.BytesIO()
    workbook.save(out)
    out.seek(0)
    return StreamingResponse(
        out,
        media_type=("application/vnd.openxmlformats-officedocument"
                    ".spreadsheetml.sheet"),
        headers={"Content-Disposition":
                 _appraisal_filename(db, project_id, "xlsx")})
