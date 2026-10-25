# SPDX-FileCopyrightText: 2026 Université Gustave Eiffel
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Import an extraction table filled on another platform (xlsx).

Flow: the admin uploads the workbook, gets a proposed header mapping
(metadata columns vs extraction columns, matched by name against the
project's grid) plus row-matching stats, adjusts it in the UI, then applies.
Rows are matched to existing records by DOI, then by normalized title;
unmatched rows can be created as new records. Every record touched by the
import is flagged `ft_forced_included` so it enters the extraction corpus
without a local full-text screening round (the review was screened elsewhere).
Imported cells are stored as the importing user's validated human values —
empty cells are skipped so the team can finish the table in the app.
"""
import io
import json
import re

from sqlalchemy.orm import Session

from app.models import ExtractionColumn, ExtractionValue, Record
from app.services.connectors.jobs import (
    archive_raw, get_or_create_source, store_import,
)
from app.services.extraction import service as ext
from app.services.extraction import studies as ext_studies
from app.services.ris.parse import normalize_doi


class ImportTableError(Exception):
    pass


META_ROLES = ("title", "doi", "authors", "year", "journal", "abstract")
# Optional: which study of the article a row describes (chantier 9). Without
# it, several rows matching the same article become its successive studies.
STUDY_ROLE = "study"

# Header aliases (lowercase) -> metadata role. Includes our own export's
# headers and common French spellings, so round-tripping an export works.
_META_ALIASES: dict[str, str] = {
    "title": "title", "article title": "title", "document title": "title",
    "titre": "title", "ti": "title",
    "doi": "doi",
    "authors": "authors", "author": "authors", "auteurs": "authors",
    "auteur": "authors", "author full names": "authors",
    "year": "year", "publication year": "year", "année": "year", "annee": "year",
    "journal": "journal", "source title": "journal", "revue": "journal",
    "journal title": "journal",
    "abstract": "abstract", "résumé": "abstract", "resume": "abstract",
    "study": STUDY_ROLE, "étude": STUDY_ROLE, "etude": STUDY_ROLE,
    "study label": STUDY_ROLE, "study name": STUDY_ROLE,
}
# Headers from our own export that carry no importable cell value.
_IGNORED_ALIASES = {"record id", "first author", "study design", "jbi score",
                    "id", "premier auteur"}

MAX_ROWS = 5000
MAX_WARNINGS = 30
# An .xlsx is a zip whose sheets are inflated to be read: five times the
# largest upload (20 MB) — far beyond 5,000 rows — so only a zip bomb is refused.
MAX_XLSX_UNPACKED_BYTES = 100 * 1024 * 1024


def parse_xlsx(data: bytes) -> tuple[list[str], list[list]]:
    """First sheet of the workbook -> (header row, data rows as raw values)."""
    from openpyxl import load_workbook

    from app.core.zipguard import ZipTooLarge, check_unpacked_size
    try:
        check_unpacked_size(data, MAX_XLSX_UNPACKED_BYTES)
    except ZipTooLarge as e:
        raise ImportTableError(f"Could not read this workbook: {e}")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True,
                                 data_only=True)
    except Exception:
        raise ImportTableError("Could not read this file as an .xlsx workbook")
    sheet = workbook.worksheets[0]
    headers: list[str] = []
    rows: list[list] = []
    for row in sheet.iter_rows(values_only=True):
        if not headers:
            if any(v is not None and str(v).strip() for v in row):
                headers = [str(v).strip() if v is not None else ""
                           for v in row]
            continue
        if any(v is not None and str(v).strip() != "" for v in row):
            rows.append(list(row[:len(headers)]))
        if len(rows) > MAX_ROWS:
            raise ImportTableError(
                f"Too many rows (max {MAX_ROWS} per import)")
    workbook.close()
    if not headers:
        raise ImportTableError("The first sheet has no header row")
    return headers, rows


def suggest_mapping(db: Session, project_id: int,
                    headers: list[str]) -> list[dict]:
    """Propose a role per header: metadata field, existing extraction column
    (matched by name, case-insensitive), new column, or ignore."""
    columns = {c.name.strip().lower(): c for c in
               db.query(ExtractionColumn).filter_by(project_id=project_id)}
    out = []
    for i, header in enumerate(headers):
        key = header.strip().lower()
        entry = {"index": i, "header": header, "role": "column",
                 "column_id": None, "column_name": None}
        if not key or key in _IGNORED_ALIASES:
            entry["role"] = "ignore"
        elif key in _META_ALIASES:
            entry["role"] = _META_ALIASES[key]
        elif key in columns:
            entry["column_id"] = columns[key].id
            entry["column_name"] = columns[key].name
        out.append(entry)
    return out


def _norm_title(title: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(title or "").lower()).strip()


def _canonical(record: Record, by_id: dict[int, Record]) -> Record:
    """Follow the dedup pointer so values land on the canonical record."""
    if record.canonical_id and record.canonical_id in by_id:
        return by_id[record.canonical_id]
    return record


def _row_meta(row: list, mapping: list[dict]) -> dict:
    meta = {}
    for m in mapping:
        if m["role"] in META_ROLES and m["index"] < len(row):
            meta[m["role"]] = row[m["index"]]
    return meta


def parse_mapping(raw: str, headers: list[str]) -> list[dict]:
    """Validate the client-sent mapping JSON against the parsed headers."""
    try:
        entries = json.loads(raw)
        assert isinstance(entries, list)
    except (ValueError, AssertionError):
        raise ImportTableError("Invalid mapping payload")
    valid_roles = set(META_ROLES) | {STUDY_ROLE, "column", "ignore"}
    mapping = []
    for e in entries:
        if not isinstance(e, dict):
            raise ImportTableError("Invalid mapping payload")
        index, role = e.get("index"), e.get("role")
        if not isinstance(index, int) or not 0 <= index < len(headers) \
                or role not in valid_roles:
            raise ImportTableError("Invalid mapping payload")
        column_id = e.get("column_id")
        mapping.append({"index": index, "header": headers[index],
                        "role": role,
                        "column_id": column_id
                        if isinstance(column_id, int) else None})
    roles = [m["role"] for m in mapping]
    for role in META_ROLES + (STUDY_ROLE,):
        if roles.count(role) > 1:
            raise ImportTableError(f"'{role}' is mapped to several columns")
    if "title" not in roles and "doi" not in roles:
        raise ImportTableError(
            "Map at least a Title or a DOI column so rows can be matched "
            "to articles")
    return mapping


def match_rows(db: Session, project_id: int, rows: list[list],
               mapping: list[dict]) -> tuple[list[Record | None], list[dict]]:
    """One entry per data row: the matched (canonical) record, or None."""
    records = db.query(Record).filter_by(project_id=project_id).all()
    by_id = {r.id: r for r in records}
    by_doi: dict[str, Record] = {}
    by_title: dict[str, Record] = {}
    for r in records:
        if r.doi and r.doi not in by_doi:
            by_doi[r.doi] = r
        key = _norm_title(r.title)
        if key and key not in by_title:
            by_title[key] = r
    matches: list[Record | None] = []
    metas: list[dict] = []
    for row in rows:
        meta = _row_meta(row, mapping)
        metas.append(meta)
        record = None
        doi = normalize_doi(meta.get("doi"))
        if doi:
            record = by_doi.get(doi)
        if record is None:
            key = _norm_title(meta.get("title") or "")
            if key:
                record = by_title.get(key)
        matches.append(_canonical(record, by_id) if record else None)
    return matches, metas


def preview(db: Session, project_id: int, data: bytes) -> dict:
    headers, rows = parse_xlsx(data)
    mapping = suggest_mapping(db, project_id, headers)
    included = set(ext.eligible_ids(db, project_id))
    matches, metas = match_rows(db, project_id, rows, mapping)
    unmatched_titles = [str(m.get("title") or m.get("doi") or "?")[:120]
                        for m, r in zip(metas, matches) if r is None][:5]
    return {
        "headers": mapping,
        "n_rows": len(rows),
        "matched": sum(1 for r in matches if r is not None),
        "matched_not_included": sum(1 for r in matches
                                    if r is not None and r.id not in included),
        "unmatched": sum(1 for r in matches if r is None),
        "unmatched_sample": unmatched_titles,
    }


def apply(db: Session, project_id: int, user_id: int, filename: str,
          data: bytes, mapping_raw: str, create_missing: bool,
          source_name: str) -> dict:
    headers, rows = parse_xlsx(data)
    mapping = parse_mapping(mapping_raw, headers)

    # 1. Resolve extraction columns: reuse mapped/name-matched ones, create
    # the rest as text columns named after their header.
    existing = {c.name.strip().lower(): c for c in
                db.query(ExtractionColumn).filter_by(project_id=project_id)}
    by_id = {c.id: c for c in existing.values()}
    position = len(existing)
    columns: dict[int, ExtractionColumn] = {}  # header index -> column
    created_columns = 0
    for m in mapping:
        if m["role"] != "column":
            continue
        column = by_id.get(m["column_id"]) if m["column_id"] else None
        if column is None:
            column = existing.get(m["header"].strip().lower())
        if column is None:
            if not m["header"].strip():
                continue
            column = ExtractionColumn(
                project_id=project_id, name=m["header"].strip()[:255],
                col_type="text", description="", options=[],
                position=position)
            db.add(column)
            db.flush()
            existing[column.name.strip().lower()] = column
            position += 1
            created_columns += 1
        columns[m["index"]] = column

    # 2. Match rows; create the missing records under a dedicated source.
    matches, metas = match_rows(db, project_id, rows, mapping)
    warnings: list[str] = []
    created_records = 0
    if create_missing and any(r is None for r in matches):
        new_metas = []
        seen_new: set[str] = set()
        for meta, record in zip(metas, matches):
            if record is not None:
                continue
            title = " ".join(str(meta.get("title") or "").split())
            doi = normalize_doi(meta.get("doi"))
            if not title and not doi:
                continue
            # several rows for one unknown article (its studies, chantier 9)
            # create that article once
            key = doi or _norm_title(title)
            if key in seen_new:
                continue
            seen_new.add(key)
            authors = [a.strip() for a in
                       re.split(r"[;|]", str(meta.get("authors") or ""))
                       if a.strip()]
            year = re.search(r"(1[89]\d{2}|20\d{2})",
                             str(meta.get("year") or ""))
            new_metas.append({
                "doi": doi, "title": title, "authors": authors,
                "year": int(year.group(1)) if year else None,
                "journal": " ".join(str(meta.get("journal") or "").split()),
                "abstract": str(meta.get("abstract") or "").strip(),
            })
        if new_metas:
            archived = archive_raw(project_id, filename, data)
            source = get_or_create_source(db, project_id, "manual",
                                          source_name.strip()[:100])
            imp = store_import(db, project_id=project_id, source_id=source.id,
                               equation_id=None, filename=archived,
                               records=new_metas, user_id=user_id)
            db.flush()
            fresh = db.query(Record).filter_by(import_id=imp.id).all()
            for r in fresh:
                r.ft_forced_included = True
            created_records = len(fresh)
            # re-match so the new records receive their cell values
            matches, metas = match_rows(db, project_id, rows, mapping)

    # 3. Force-include matched records that are not FT-included here.
    included = set(ext.eligible_ids(db, project_id))
    forced = 0
    for record in matches:
        if record is not None and record.id not in included \
                and not record.ft_forced_included:
            record.ft_forced_included = True
            included.add(record.id)
            forced += 1

    # 4. Write the cells as the importer's validated values (upsert; empty
    # cells are skipped so the table can be completed in the app). A row goes
    # to the study its Study column names (created when new); without that
    # column, the k-th row of an article is its k-th study.
    touched_ids = {record.id for record in matches if record is not None}
    values = {(v.study_id, v.column_id): v for v in
              db.query(ExtractionValue).filter(
                  ExtractionValue.record_id.in_(touched_ids),
                  ExtractionValue.user_id == user_id)} if touched_ids else {}
    study_index = next((m["index"] for m in mapping if m["role"] == STUDY_ROLE),
                       None)
    seen_rows: dict[int, int] = {}
    studies_created = 0
    cells_written = 0
    for row, record in zip(rows, matches):
        if record is None:
            continue
        label = (str(row[study_index]).strip()
                 if study_index is not None and study_index < len(row)
                 and row[study_index] is not None else "")
        studies = ext_studies.studies_for_records(db, [record.id], user_id)[record.id]
        if label:
            study = ext_studies.find_by_label(db, record, label)
            if study is None and len(studies) == 1 \
                    and studies[0].label == ext_studies.DEFAULT_LABEL \
                    and not ext_studies.has_content(db, studies[0]):
                # the untouched default study takes the sheet's name rather
                # than sitting empty next to it
                study = studies[0]
                ext_studies.rename_study(db, study, label)
            if study is None:
                study = ext_studies.add_study(db, record, label, user_id)
                studies_created += 1
        else:
            k = seen_rows.get(record.id, 0)
            if k < len(studies):
                study = studies[k]
            else:
                study = ext_studies.add_study(db, record, f"Study {k + 1}", user_id)
                studies_created += 1
        seen_rows[record.id] = seen_rows.get(record.id, 0) + 1
        for index, column in columns.items():
            raw = row[index] if index < len(row) else None
            if raw is None or str(raw).strip() == "":
                continue
            if column.col_type == "multi" and isinstance(raw, str):
                raw = [s.strip() for s in raw.split(";") if s.strip()]
            try:
                stored = ext.coerce_value(column, raw)
            except ValueError as e:
                if len(warnings) < MAX_WARNINGS:
                    warnings.append(
                        f"'{column.name}' — {str(record.title)[:60]}: {e}")
                continue
            value = values.get((study.id, column.id))
            if value is None:
                value = ExtractionValue(record_id=record.id, study_id=study.id,
                                        column_id=column.id, user_id=user_id)
                db.add(value)
                values[(study.id, column.id)] = value
            value.value = stored
            value.origin = "human"
            value.validated = True
            value.imported_from = f"xlsx:{filename[:200]}"
            cells_written += 1

    # 5. Keep the allocation in step when extraction already started.
    config = ext.get_config(db, project_id)
    if config is not None and config.status == "active":
        ext.generate_assignments(db, config)

    return {
        "n_rows": len(rows),
        "matched": sum(1 for r in matches if r is not None),
        "records_created": created_records,
        "records_force_included": forced + created_records,
        "columns_created": created_columns,
        "studies_created": studies_created,
        "cells_written": cells_written,
        "skipped_rows": sum(1 for r in matches if r is None),
        "warnings": warnings,
    }
