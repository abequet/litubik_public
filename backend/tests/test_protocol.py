import io
import time

import pytest

from app.services.llm.base import LLMResult
from app.services.protocol import (
    ProtocolFileError,
    extract_text,
    normalize_criteria,
    render_protocol_md,
)

STRUCTURED = {
    "title": "Effects of X on Y",
    "background": "Some theoretical background.",
    "research_questions": ["Does X improve Y?"],
    "objectives": ["Synthesize evidence on X"],
    "population": "Adults 18+",
    "scope": "Empirical studies, 2010-2025, EN/FR",
    "inclusion_criteria": [
        {"label": "Adult population", "description": "Participants aged 18 or older"},
        {"label": "Empirical study", "description": "Reports primary empirical data"},
    ],
    "exclusion_criteria": [
        {"label": "Wrong publication type", "description": "Editorials, letters, abstracts only"},
    ],
}


# ---- text extraction --------------------------------------------------------

def test_extract_txt_and_md():
    # runs of blank lines are collapsed to at most two
    assert extract_text("protocol.txt", "Line one\n\n\n\n\nLine two".encode()) == \
        "Line one\n\n\nLine two"
    assert "Protocole" in extract_text("proto.md", "# Protocole\ncontenu".encode())


def test_extract_txt_latin1_fallback():
    assert "éligibilité" in extract_text("p.txt", "critères d'éligibilité".encode("latin-1"))


def test_extract_pdf():
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Inclusion criteria: adults only")
    data = doc.tobytes()
    doc.close()
    assert "adults only" in extract_text("protocol.pdf", data)


def test_extract_docx():
    import docx
    document = docx.Document()
    document.add_paragraph("Research question: does X improve Y?")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "E1"
    table.rows[0].cells[1].text = "Not empirical"
    buf = io.BytesIO()
    document.save(buf)
    text = extract_text("protocol.docx", buf.getvalue())
    assert "does X improve Y" in text
    assert "Not empirical" in text


def test_extract_unsupported_extension():
    with pytest.raises(ProtocolFileError, match="Unsupported"):
        extract_text("protocol.rtf", b"whatever")


def test_extract_empty_pdf_raises():
    import pymupdf
    doc = pymupdf.open()
    doc.new_page()
    data = doc.tobytes()
    doc.close()
    with pytest.raises(ProtocolFileError, match="No text"):
        extract_text("blank.pdf", data)


def test_extract_corrupt_pdf_raises():
    with pytest.raises(ProtocolFileError):
        extract_text("bad.pdf", b"this is not a pdf")


# ---- structuring ------------------------------------------------------------

def test_normalize_criteria_codes():
    criteria = normalize_criteria(STRUCTURED)
    assert [c["code"] for c in criteria] == ["I1", "I2", "E1"]
    assert criteria[0]["type"] == "include"
    assert criteria[2]["type"] == "exclude"
    assert criteria[2]["label"] == "Wrong publication type"


def test_normalize_criteria_skips_empty_items():
    structured = {"inclusion_criteria": [{"label": "", "description": ""}],
                  "exclusion_criteria": [{"label": "Kept", "description": "d"}]}
    criteria = normalize_criteria(structured)
    assert [c["code"] for c in criteria] == ["E1"]


def test_render_protocol_md():
    criteria = normalize_criteria(STRUCTURED)
    md = render_protocol_md(STRUCTURED, criteria)
    assert md.startswith("# Effects of X on Y")
    assert "## Research questions" in md
    assert "1. Does X improve Y?" in md
    assert "- **I2 — Empirical study**: Reports primary empirical data" in md
    assert "- **E1 — Wrong publication type**" in md
    assert "**Population:** Adults 18+" in md
    # nothing extracted -> no empty section
    assert "## Search concepts" not in md


def test_render_protocol_md_keeps_the_search_concepts():
    """protocol_md is the only durable channel: what the protocol says about
    its own concepts and keywords has to survive the rendering, or the equation
    builder never sees it."""
    structured = {**STRUCTURED, "search_concepts": [
        {"name": "Population", "keywords": ["older adults", "elderly"]},
        {"name": "Outcome", "keywords": []},
    ], "search_strategy_notes": "PubMed and Scopus, 2010 onwards."}
    md = render_protocol_md(structured, normalize_criteria(structured))
    assert "## Search concepts" in md
    assert "- **Population**: older adults, elderly" in md
    assert "- **Outcome**: —" in md
    assert "**Search strategy:** PubMed and Scopus, 2010 onwards." in md
    assert md.index("## Search concepts") < md.index("## Inclusion criteria")


# ---- question framing -------------------------------------------------------

def test_extraction_schema_follows_the_project_framework():
    from app.services.protocol import extraction_schema

    picos = extraction_schema("picos")["properties"]["framing"]["properties"]
    assert "population" in picos and "study_design" in picos

    spider = extraction_schema("spider")["properties"]["framing"]["properties"]
    assert "phenomenon_of_interest" in spider
    assert "intervention" not in spider  # SPIDER has no intervention slot


def test_extraction_schema_never_requires_a_framing_field():
    """A protocol that says nothing about its population must still extract."""
    from app.services.protocol import EXTRACTION_SCHEMA, extraction_schema

    for key in ("picos", "spider", "free"):
        required = extraction_schema(key)["required"]
        assert "framing" not in required
        assert required == EXTRACTION_SCHEMA["required"]


def test_render_protocol_md_renders_the_framework_slots():
    structured = {**STRUCTURED, "framing": {
        "sample": "Ambulance crews",
        "phenomenon_of_interest": "Acute stress while driving under blue light",
        "research_type": "Qualitative and mixed methods",
    }}
    md = render_protocol_md(structured, normalize_criteria(structured), "spider")
    assert "## Question framing (SPIDER)" in md
    assert "**Sample:** Ambulance crews" in md
    assert "**Phenomenon of interest:** Acute stress while driving under blue light" in md
    # a slot the document says nothing about produces no empty line
    assert "**Design:**" not in md
    assert "**Scope:** Empirical studies, 2010-2025, EN/FR" in md


def test_render_protocol_md_keeps_the_historical_heading_for_pico():
    structured = {**STRUCTURED, "framing": {"population": "Adults 18+"}}
    md = render_protocol_md(structured, normalize_criteria(structured), "picos")
    assert "## Population & scope" in md
    assert "**Population:** Adults 18+" in md


def test_render_protocol_md_falls_back_to_the_pre_framework_shape():
    """Protocols extracted before framings existed — and models that answer in
    the old form — must still render."""
    md = render_protocol_md(STRUCTURED, normalize_criteria(STRUCTURED), "spider")
    assert "**Population:** Adults 18+" in md


# ---- API flow ---------------------------------------------------------------

@pytest.fixture
def project(user_client):
    r = user_client.post("/api/projects", json={"name": "SR demo"})
    return r.json()


def _upload(client, project_id, text="Protocol text: include adults, exclude editorials."):
    return client.post(
        f"/api/projects/{project_id}/protocols/upload",
        files={"file": ("protocol.txt", text.encode(), "text/plain")},
    )


def test_upload_edit_validate_flow(user_client, project):
    pid = project["id"]

    r = _upload(user_client, pid)
    assert r.status_code == 200, r.text
    protocol = r.json()
    assert protocol["version"] == 1
    assert protocol["source_text_chars"] > 0
    assert protocol["has_content"] is False

    # nothing to validate yet
    r = user_client.post(f"/api/projects/{pid}/protocols/{protocol['id']}/validate")
    assert r.status_code == 400

    # manual edit (protocol_md + criteria) works without any LLM
    r = user_client.put(f"/api/projects/{pid}/protocols/{protocol['id']}", json={
        "protocol_md": "# My protocol\n\nEdited by hand.",
        "criteria": [
            {"code": "i1", "type": "include", "label": "Adults", "description": "18+"},
            {"code": "E1", "type": "exclude", "label": "Editorials", "description": ""},
        ],
    })
    assert r.status_code == 200
    body = r.json()
    assert body["criteria"][0]["code"] == "I1"  # codes normalized to uppercase

    # duplicate codes rejected
    r = user_client.put(f"/api/projects/{pid}/protocols/{protocol['id']}", json={
        "protocol_md": "x",
        "criteria": [
            {"code": "I1", "type": "include", "label": "a", "description": ""},
            {"code": "I1", "type": "include", "label": "b", "description": ""},
        ],
    })
    assert r.status_code == 400

    # validate freezes the version
    r = user_client.post(f"/api/projects/{pid}/protocols/{protocol['id']}/validate")
    assert r.status_code == 200
    assert r.json()["validated_at"] is not None

    r = user_client.put(f"/api/projects/{pid}/protocols/{protocol['id']}", json={
        "protocol_md": "nope", "criteria": []})
    assert r.status_code == 400
    r = user_client.delete(f"/api/projects/{pid}/protocols/{protocol['id']}")
    assert r.status_code == 400

    # new upload becomes version 2; list shows both
    r = _upload(user_client, pid, "Amended protocol")
    assert r.json()["version"] == 2
    versions = user_client.get(f"/api/projects/{pid}/protocols").json()
    assert [v["version"] for v in versions] == [2, 1]

    # unvalidate reopens editing
    r = user_client.post(f"/api/projects/{pid}/protocols/{protocol['id']}/unvalidate")
    assert r.status_code == 200
    assert r.json()["validated_at"] is None


def test_paste_protocol_creates_version(user_client, project):
    """Writing/pasting protocol text (no file) creates a version just like an
    upload, with the pasted text as the source text."""
    pid = project["id"]
    text = "Objective: assess X. Include adults. Exclude editorials."
    r = user_client.post(f"/api/projects/{pid}/protocols/paste",
                         json={"text": text, "name": "Registered protocol"})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["source_text_chars"] == len(text)
    assert body["source_filename"] == "Registered protocol.txt"
    # it shows up as a normal version, editable and generatable
    versions = user_client.get(f"/api/projects/{pid}/protocols").json()
    assert any(v["id"] == body["id"] for v in versions)
    # empty paste is rejected
    assert user_client.post(f"/api/projects/{pid}/protocols/paste",
                            json={"text": "   "}).status_code == 400


def test_upload_rejects_bad_type(user_client, project):
    r = user_client.post(
        f"/api/projects/{project['id']}/protocols/upload",
        files={"file": ("protocol.rtf", b"x", "application/rtf")},
    )
    assert r.status_code == 400


def test_reviewer_cannot_upload_or_validate(user_client, project):
    pid = project["id"]
    # second account, added as reviewer
    user_client.post("/api/auth/register", json={
        "email": "rev@example.com", "password": "password123", "display_name": "Rev"})
    # registering switched the session to the reviewer — go back to the admin
    user_client.post("/api/auth/login", json={
        "email": "admin@example.com", "password": "password123"})
    user_client.post(f"/api/projects/{pid}/members",
                     json={"email": "rev@example.com", "role": "reviewer"})
    r = _upload(user_client, pid)
    protocol_id = r.json()["id"]

    user_client.post("/api/auth/login", json={
        "email": "rev@example.com", "password": "password123"})
    assert _upload(user_client, pid).status_code == 403
    r = user_client.post(f"/api/projects/{pid}/protocols/{protocol_id}/validate")
    assert r.status_code == 403
    # but reviewers can read
    assert user_client.get(f"/api/projects/{pid}/protocols").status_code == 200


def test_generate_job_fills_protocol(user_client, project, monkeypatch):
    pid = project["id"]
    protocol_id = _upload(user_client, pid).json()["id"]
    # a stored credential is required for resolution
    user_client.put("/api/llm/credentials", json={
        "provider": "anthropic", "api_key": "sk-ant-fake-key-123"})

    class FakeProvider:
        model = "claude-fake"
        name = "anthropic"

        def generate_json(self, messages, json_schema, system=None, max_tokens=4096):
            assert "protocol_document" in messages[0]["content"]
            return LLMResult(text="{}", tokens_input=1200, tokens_output=300,
                             data=STRUCTURED)

    monkeypatch.setattr("app.services.protocol.jobs.open_provider_chain",
                        lambda *a, **kw: FakeProvider())

    r = user_client.post(
        f"/api/projects/{pid}/protocols/{protocol_id}/generate", json={})
    assert r.status_code == 200, r.text
    job_id = r.json()["job_id"]

    for _ in range(100):
        job = user_client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "failed"):
            break
        time.sleep(0.05)
    assert job["status"] == "done", job.get("error")
    assert job["result"]["n_criteria"] == 3
    # per-call cost is reported (1200 in / 300 out at anthropic sonnet rates)
    assert job["result"]["cost"]["cost_usd"] > 0
    assert job["result"]["cost"]["tokens_in"] == 1200

    protocol = user_client.get(f"/api/projects/{pid}/protocols/{protocol_id}").json()
    assert protocol["has_content"] is True
    assert "# Effects of X on Y" in protocol["protocol_md"]
    assert [c["code"] for c in protocol["criteria"]] == ["I1", "I2", "E1"]
    assert protocol["generated_with"] == "anthropic:claude-fake"
