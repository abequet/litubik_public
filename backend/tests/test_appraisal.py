"""Critical appraisal: instrument registry, two-assessor mode, consensus.

What is under test is CREST_Triage domain 4 — that an instrument fits the
studies, and that a second assessor's disagreement is recorded rather than
overwritten. The historical behaviour (any member appraises with JBI, no
disagreement computed) must survive untouched for a project that never opens
the setting: that is the non-regression this file guards.
"""
from tests.test_extraction import ext_project, team, _login, _await_job  # noqa: F401


def _appraise(client, pid, rid, answers, overall="include"):
    return client.put(f"/api/projects/{pid}/extraction/records/{rid}/jbi",
                      json={"answers": answers, "overall": overall})


# -- the registry -------------------------------------------------------------

def test_the_shipped_instruments_are_jbi_and_mmat_and_cover_their_designs():
    from app.services.appraisal import instruments
    assert set(instruments.BUILTIN) == {"jbi", "mmat"}
    for inst in instruments.BUILTIN.values():
        assert inst.builtin and inst.licence
        # every checklist keys into the shared design vocabulary
        assert set(inst.checklists) <= set(instruments.DESIGNS)
        for design, checklist in inst.checklists.items():
            qids = instruments.question_ids(inst, design)
            assert qids, design
            assert len(set(qids)) == len(qids), (inst.key, design)
            assert set(checklist["options"]) <= set(instruments.OPTION_SETS["ynun"])
    jbi, mmat = instruments.BUILTIN["jbi"], instruments.BUILTIN["mmat"]
    assert jbi.overall and jbi.score
    # the MMAT guide discourages any score and has no include/exclude verdict
    assert not mmat.overall and not mmat.score


def test_mmat_is_transcribed_as_published():
    """Two screening questions, five categories of five criteria, answers
    yes / no / can't tell; a mixed-methods study is rated on 5.1–5.5 and on
    its component criteria; reviews are out of the tool's scope."""
    from app.services.appraisal import instruments, mmat
    inst = instruments.BUILTIN["mmat"]
    assert "systematic_review" not in inst.checklists
    assert "other" not in inst.checklists
    for design in ("qualitative", "rct", "cohort", "cross_sectional"):
        qids = instruments.question_ids(inst, design)
        assert qids[:2] == ["S1", "S2"]
        assert len(qids) == 7
        # the screening questions are optional, the category's five are not
        assert len(instruments.required_question_ids(inst, design)) == 5
    for design in ("quasi_experimental", "cohort", "case_control"):
        assert instruments.required_question_ids(inst, design) == [
            "3.1", "3.2", "3.3", "3.4", "3.5"]
    mixed = instruments.question_ids(inst, "mixed_methods")
    assert set(mixed) == {"S1", "S2"} | {
        f"{c}.{i}" for c in range(1, 6) for i in range(1, 6)}
    assert instruments.required_question_ids(inst, "mixed_methods") == [
        "5.1", "5.2", "5.3", "5.4", "5.5"]
    checklist = inst.checklists["rct"]
    assert checklist["options"] == ["yes", "no", "unclear"]
    assert checklist["option_labels"] == {"unclear": "Can't tell"}
    assert mmat.CATEGORY_DESIGNS["3"] == ["quasi_experimental", "cohort",
                                          "case_control"]


def test_the_design_vocabulary_is_wider_than_any_one_instrument():
    """A study's design does not change with the tool used to appraise it —
    mixed-methods studies exist whether or not JBI has a checklist for them."""
    from app.services.appraisal import instruments
    assert "mixed_methods" in instruments.DESIGNS
    assert "mixed_methods" not in instruments.BUILTIN["jbi"].checklists


def test_a_project_declares_its_own_instrument(user_client, ext_project):
    pid = ext_project["pid"]
    r = user_client.put(f"/api/projects/{pid}/appraisal/instruments/lab-tool",
                        json={"key": "lab-tool", "label": "Our lab checklist",
                              "citation": "Internal, 2026",
                              "checklists": {"mixed_methods": {
                                  "options": "ynu",
                                  "domains": [{"label": "Sampling", "questions": [
                                      {"id": "s1", "text": "Was the sample described?"}]}]}}})
    assert r.status_code == 200, r.text
    rows = user_client.get(f"/api/projects/{pid}/appraisal/instruments").json()
    keys = [i["key"] for i in rows["instruments"]]
    assert keys == ["jbi", "mmat", "lab-tool"]  # built-in first
    custom = next(i for i in rows["instruments"] if i["key"] == "lab-tool")
    assert custom["builtin"] is False
    assert custom["designs"] == ["mixed_methods"]


def test_a_declared_instrument_cannot_shadow_a_builtin_or_be_empty(
        user_client, ext_project):
    pid = ext_project["pid"]
    r = user_client.put(f"/api/projects/{pid}/appraisal/instruments/jbi",
                        json={"key": "jbi", "label": "Mine", "checklists": {}})
    assert r.status_code == 400
    # a checklist with no usable question would render as an empty form
    r = user_client.put(f"/api/projects/{pid}/appraisal/instruments/x",
                        json={"key": "x", "label": "X", "checklists": {
                            "rct": {"options": "ynu", "domains": []}}})
    assert r.status_code == 400
    r = user_client.put(f"/api/projects/{pid}/appraisal/instruments/Bad Key",
                        json={"key": "Bad Key", "label": "X", "checklists": {}})
    assert r.status_code == 400


def test_the_instrument_in_use_cannot_be_deleted(user_client, ext_project):
    pid = ext_project["pid"]
    user_client.put(f"/api/projects/{pid}/appraisal/instruments/lab-tool",
                    json={"key": "lab-tool", "label": "Lab", "checklists": {
                        "rct": {"options": "ynu", "domains": [
                            {"label": "D", "questions": [
                                {"id": "q1", "text": "Q?"}]}]}}})
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "lab-tool",
                          "n_assessors": 1})
    r = user_client.delete(f"/api/projects/{pid}/appraisal/instruments/lab-tool")
    assert r.status_code == 400
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "jbi"})
    assert user_client.delete(
        f"/api/projects/{pid}/appraisal/instruments/lab-tool").status_code == 200


def test_switching_instrument_changes_the_checklist_served(
        user_client, ext_project):
    pid = ext_project["pid"]
    user_client.put(f"/api/projects/{pid}/appraisal/instruments/lab-tool",
                    json={"key": "lab-tool", "label": "Lab", "checklists": {
                        "rct": {"options": "ynu", "domains": [
                            {"label": "D", "questions": [
                                {"id": "z1", "text": "Only question"}]}]}}})
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "lab-tool"})
    served = user_client.get(
        f"/api/projects/{pid}/extraction/jbi/checklists").json()
    assert served["instrument"]["key"] == "lab-tool"
    assert [c["key"] for c in served["checklists"]] == ["rct"]

    # and the appraisal form now only accepts that instrument's question ids
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "rct"})
    _appraise(user_client, pid, rid, {"z1": {"answer": "yes", "comment": ""}})
    from app.core.db import SessionLocal
    from app.models import JbiAppraisal
    with SessionLocal() as db:
        row = db.query(JbiAppraisal).filter_by(record_id=rid).first()
    assert row.instrument == "lab-tool"
    assert set(row.answers) == {"z1"}


def test_an_unknown_instrument_key_falls_back_to_jbi(user_client, ext_project):
    """A snapshot naming an instrument this version does not know must not
    lock the team out of its own appraisals."""
    from app.core.db import SessionLocal
    from app.models import AppraisalConfig
    from app.services.appraisal import core
    pid = ext_project["pid"]
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "jbi"})
    with SessionLocal() as db:
        db.query(AppraisalConfig).filter_by(project_id=pid).update(
            {"instrument": "from-the-future"})
        db.commit()
        assert core.instrument_of(db, pid).key == "jbi"


# -- staying optional ---------------------------------------------------------

def test_a_project_that_never_opens_the_setting_behaves_as_before(
        user_client, ext_project):
    pid = ext_project["pid"]
    config = user_client.get(f"/api/projects/{pid}/appraisal/config").json()
    assert config["status"] == "draft"
    assert config["dual"] is False
    assert config["instrument"] == "jbi"
    # no dual mode -> no conflict list, whatever the assessors said
    conflicts = user_client.get(f"/api/projects/{pid}/appraisal/conflicts").json()
    assert conflicts == {"dual": False, "conflicts": []}


def test_consensus_is_refused_when_the_project_is_not_dual(
        user_client, ext_project):
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    r = user_client.put(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
        json={"overall": "include"})
    assert r.status_code == 400


# -- two assessors ------------------------------------------------------------

def _dual_project(user_client, team, ext_project, scope="overall"):
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "jbi",
                          "n_assessors": 2, "conflict_scope": scope,
                          "assessor_ids": [team["admin"]["id"],
                                           team["reviewer"]["id"]]})
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "rct"})
    return pid, rid


def test_two_assessors_who_agree_produce_no_conflict(
        user_client, team, ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    _appraise(user_client, pid, rid, {"q1": {"answer": "yes", "comment": ""}},
              overall="include")
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {"q1": {"answer": "no", "comment": ""}},
              overall="include")
    _login(user_client, "admin")
    conflicts = user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()
    assert conflicts["dual"] is True
    # they differ on q1 but agree overall — under "overall" scope that is not
    # a disagreement to arbitrate
    assert conflicts["conflicts"] == []


def test_per_question_scope_sees_what_the_overall_scope_ignores(
        user_client, team, ext_project):
    pid, rid = _dual_project(user_client, team, ext_project,
                             scope="per_question")
    _appraise(user_client, pid, rid, {"q1": {"answer": "yes", "comment": ""}})
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {"q1": {"answer": "no", "comment": ""}})
    _login(user_client, "admin")
    rows = user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()["conflicts"]
    assert [c["conflicts"] for c in rows] == [["q1"]]


def test_a_blank_answer_is_not_a_disagreement(user_client, team, ext_project):
    """An assessor who has not answered yet has not contradicted anyone."""
    pid, rid = _dual_project(user_client, team, ext_project,
                             scope="per_question")
    _appraise(user_client, pid, rid, {"q1": {"answer": "yes", "comment": ""}})
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {"q1": {"answer": "", "comment": ""}},
              overall="include")
    _login(user_client, "admin")
    rows = user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()["conflicts"]
    assert rows == []


def test_a_disagreement_is_resolved_without_touching_what_each_one_said(
        user_client, team, ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    _appraise(user_client, pid, rid, {"q1": {"answer": "yes", "comment": ""}},
              overall="include")
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {"q1": {"answer": "no", "comment": ""}},
              overall="exclude")
    _login(user_client, "admin")

    rows = user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()["conflicts"]
    assert [c["conflicts"] for c in rows] == [["overall"]]

    view = user_client.get(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus").json()
    assert len(view["appraisals"]) == 2
    assert view["can_resolve"] is True

    r = user_client.put(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
        json={"answers": {"q1": {"answer": "unclear", "comment": "reread p.4"}},
              "overall": "include", "comment": "agreed after discussion"})
    assert r.status_code == 200

    # the individual appraisals are untouched — the disagreement stays visible
    view = user_client.get(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus").json()
    assert {a["overall"] for a in view["appraisals"]} == {"include", "exclude"}
    assert view["consensus"]["overall"] == "include"
    assert user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()["conflicts"] == []


def test_reopening_puts_the_disagreement_back(user_client, team, ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    _appraise(user_client, pid, rid, {}, overall="include")
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {}, overall="exclude")
    _login(user_client, "admin")
    user_client.put(f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
                    json={"overall": "include"})
    user_client.delete(f"/api/projects/{pid}/appraisal/records/{rid}/consensus")
    rows = user_client.get(
        f"/api/projects/{pid}/appraisal/conflicts").json()["conflicts"]
    assert [c["record_id"] for c in rows] == [rid]


def test_arbiter_mode_keeps_a_plain_assessor_out(user_client, team, ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "active", "instrument": "jbi",
                          "n_assessors": 2, "conflict_resolution": "arbiter",
                          "assessor_ids": [team["admin"]["id"],
                                           team["reviewer"]["id"]]})
    _appraise(user_client, pid, rid, {}, overall="include")
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid, {}, overall="exclude")
    r = user_client.put(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
        json={"overall": "include"})
    assert r.status_code == 403
    _login(user_client, "admin")
    assert user_client.put(
        f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
        json={"overall": "include"}).status_code == 200


# -- what the aggregates and the dossier report -------------------------------

def test_the_consensus_is_what_the_results_aggregate(user_client, team,
                                                     ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    _appraise(user_client, pid, rid,
              {f"q{i}": {"answer": "yes", "comment": ""} for i in range(1, 14)})
    _login(user_client, "reviewer")
    _appraise(user_client, pid, rid,
              {f"q{i}": {"answer": "no", "comment": ""} for i in range(1, 14)},
              overall="exclude")
    _login(user_client, "admin")
    user_client.put(f"/api/projects/{pid}/appraisal/records/{rid}/consensus",
                    json={"answers": {f"q{i}": {"answer": "yes", "comment": ""}
                                      for i in range(1, 14)},
                          "overall": "include"})
    data = user_client.get(f"/api/projects/{pid}/results").json()
    study = next(s for s in data["jbi"]["per_study"] if s["record_id"] == rid)
    assert study["source"] == "consensus"
    assert study["pct_yes"] == 1.0        # not the average of 13 yes and 13 no
    assert study["n_assessors"] == 2
    assert data["jbi"]["n_from_consensus"] == 1


def test_the_results_report_bias_domains_not_only_a_score(user_client, team,
                                                         ext_project):
    pid, rid = _dual_project(user_client, team, ext_project)
    _appraise(user_client, pid, rid, {
        "q1": {"answer": "yes", "comment": ""},
        "q2": {"answer": "no", "comment": ""},
        "q4": {"answer": "unclear", "comment": ""}})
    data = user_client.get(f"/api/projects/{pid}/results").json()["jbi"]
    block = next(b for b in data["by_domain"] if b["design"] == "rct")
    assert block["granularity"] == "domain"
    selection = next(r for r in block["rows"]
                     if r["label"].startswith("Bias related to selection"))
    assert selection["counts"] == {"yes": 1, "no": 1, "unclear": 0, "na": 0}
    assert selection["n_studies"] == 1
    # a domain nobody answered on counts no study
    retention = next(r for r in block["rows"]
                     if r["label"].startswith("Bias related to participant"))
    assert retention["n_studies"] == 0


def test_a_flat_checklist_falls_back_to_one_row_per_question(
        user_client, team, ext_project):
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "cohort"})  # single unnamed domain
    _appraise(user_client, pid, rid, {"q1": {"answer": "yes", "comment": ""}})
    data = user_client.get(f"/api/projects/{pid}/results").json()["jbi"]
    block = next(b for b in data["by_domain"] if b["design"] == "cohort")
    assert block["granularity"] == "question"
    assert block["rows"][0]["label"].startswith("Were the two groups similar")


def test_the_dossier_states_whether_appraisal_was_duplicated(
        user_client, team, ext_project):
    from app.core.db import SessionLocal
    from app.services.reporting.audit import review_dossier
    pid = ext_project["pid"]
    with SessionLocal() as db:
        block = review_dossier(db, pid)["extraction"]["appraisal"]
    assert block["dual_assessment"] is False
    assert "single assessor" in block["note"]

    _dual_project(user_client, team, ext_project)
    with SessionLocal() as db:
        block = review_dossier(db, pid)["extraction"]["appraisal"]
    assert block["dual_assessment"] is True
    assert block["independent_assessors"] == 2
    assert "consensus" in block["note"]
    assert block["instrument"].startswith("JBI")


# -- MMAT ---------------------------------------------------------------------

def test_an_mmat_appraisal_is_complete_without_an_overall_judgement(
        user_client, ext_project):
    """MMAT has no include/exclude verdict: an appraisal is done once the five
    criteria of the category are rated, screening questions or not — and a
    verdict a stale form might send is not recorded."""
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "draft", "instrument": "mmat"})
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "cohort"})
    served = user_client.get(
        f"/api/projects/{pid}/extraction/jbi/checklists").json()
    assert served["instrument"] == {
        "key": "mmat", "label": "Mixed Methods Appraisal Tool (MMAT) 2018",
        "citation": "Hong et al., 2018",
        "url": "http://mixedmethodsappraisaltoolpublic.pbworks.com",
        "builtin": True, "overall": False, "score": False}

    partial = {"3.1": {"answer": "yes", "comment": ""},
               "3.2": {"answer": "unclear", "comment": "not reported"}}
    r = _appraise(user_client, pid, rid, partial, overall="include")
    assert r.status_code == 200, r.text
    table = user_client.get(f"/api/projects/{pid}/extraction/table").json()
    assert table["instrument"]["score"] is False
    row = next(x for x in table["rows"] if x["record_id"] == rid)
    assert row["jbi"]["mine_done"] is False
    assert row["jbi"]["n_appraisals"] == 0
    assert row["jbi"]["scores"][0]["n_answered"] == 2
    assert row["jbi"]["scores"][0]["n_questions"] == 7

    full = {f"3.{i}": {"answer": "yes", "comment": ""} for i in range(1, 6)}
    _appraise(user_client, pid, rid, full, overall="include")
    table = user_client.get(f"/api/projects/{pid}/extraction/table").json()
    row = next(x for x in table["rows"] if x["record_id"] == rid)
    assert row["jbi"]["mine_done"] is True
    assert row["jbi"]["n_appraisals"] == 1
    from app.core.db import SessionLocal
    from app.models import JbiAppraisal
    with SessionLocal() as db:
        saved = db.query(JbiAppraisal).filter_by(record_id=rid).first()
    assert saved.instrument == "mmat"
    assert saved.overall == ""

    # a JBI question id is not one of this instrument: silently dropped,
    # an answer outside yes / no / can't tell refused
    r = _appraise(user_client, pid, rid,
                  {"q1": {"answer": "yes", "comment": ""},
                   "3.1": {"answer": "na", "comment": ""}})
    assert r.status_code == 400


def test_the_appraisal_export_follows_the_instrument_in_use(user_client,
                                                            ext_project):
    """The wide table has one column per criterion of the instrument, in the
    order of the published tool, spells answers its way and carries a score
    only when the instrument has one."""
    import csv
    import io
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "draft", "instrument": "mmat"})
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "qualitative"})
    _appraise(user_client, pid, rid,
              {"S1": {"answer": "yes", "comment": ""},
               "1.1": {"answer": "yes", "comment": ""},
               "1.2": {"answer": "unclear", "comment": "no detail"}},
              overall="")
    text = user_client.get(
        f"/api/projects/{pid}/extraction/export/jbi.csv").text
    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff"))))
    header = rows[0]
    assert "Score (yes)" not in header
    assert "Overall appraisal" not in header
    assert header[header.index("Instrument") + 1] == "Assessor"
    first = header.index("S1")
    assert header[first:first + 8] == ["S1", "S2", "1.1", "1.2", "1.3", "1.4",
                                       "1.5", "2.1"]
    line = dict(zip(header, rows[1]))
    assert line["Instrument"] == "mmat"
    assert line["Design"] == "qualitative"
    assert line["Answered"] == "3" and line["Out of"] == "7"
    assert line["S1"] == "yes"
    assert line["1.2"] == "Can't tell"
    # a criterion of another category is not a blank answer, it is no column
    # for this design
    assert line["2.1"] == ""

    # the details sheet of the workbook spells the answer the same way
    r = user_client.get(f"/api/projects/{pid}/extraction/export/jbi.xlsx")
    assert r.status_code == 200
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Appraisal", "Details"]
    details = list(wb["Details"].iter_rows(values_only=True))
    assert "Overall" not in details[0]
    answers = {row[details[0].index("Criterion")]: row[details[0].index("Answer")]
               for row in details[1:]}
    assert answers["1.2"] == "Can't tell"


def test_switching_instrument_hides_no_appraisal_from_the_export(
        user_client, ext_project):
    """A team that appraised under JBI then moved to MMAT keeps every answer
    in the export, under the ids they were given."""
    pid = ext_project["pid"]
    rid = ext_project["included"][0]
    user_client.put(f"/api/projects/{pid}/extraction/records/{rid}/design",
                    json={"design": "rct"})
    _appraise(user_client, pid, rid, {"q1": {"answer": "no", "comment": ""}})
    user_client.put(f"/api/projects/{pid}/appraisal/config",
                    json={"status": "draft", "instrument": "mmat"})
    import csv
    import io
    text = user_client.get(
        f"/api/projects/{pid}/extraction/export/jbi.csv").text
    rows = list(csv.reader(io.StringIO(text.lstrip("\ufeff"))))
    line = dict(zip(rows[0], rows[1]))
    assert line["Instrument"] == "jbi"
    assert line["q1"] == "no"
    assert line["2.1"] == ""
